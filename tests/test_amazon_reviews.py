import asyncio
import importlib.util
import json
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook


MODULE_PATH = Path(__file__).parents[1] / "scripts" / "amazon_reviews.py"
SPEC = importlib.util.spec_from_file_location("amazon_reviews", MODULE_PATH)
reviews = importlib.util.module_from_spec(SPEC)
sys.modules[SPEC.name] = reviews
SPEC.loader.exec_module(reviews)


class RufusParsingTests(unittest.TestCase):
    def test_public_artifacts_remove_technical_provenance(self):
        row = {
            "下载ASIN": "B000000001", "评论ID": "R1", "评论标题": "Title", "评论内容": "Body",
            "评论人": "User", "评论人链接": "https://www.amazon.com/profile/P1", "评分": 5.0,
            "评论时间": "2026-08-24", "source_url": "https://internal.example/path", "query": "seat",
            "capture_time": "now", "http_status": 200, "batch": 1,
        }
        self.assertEqual(set(reviews._public_review(row)), set(reviews.RUFUS_CSV_HEADERS))
        state = reviews._checkpoint_state({
            "asin": "B000000001", "status": "running", "stop_reason": "", "seed_complete": True,
            "top30": ["seat"], "completed_queries": ["seat"], "reviews": [row],
            "startup": {"source_url": "https://internal.example"},
            "seed": {"html": "<html>private</html>"},
            "query_responses": [{"html": "<html>private</html>", "source_url": "https://internal.example"}],
        })
        self.assertNotIn("startup", state)
        self.assertNotIn("seed", state)
        self.assertNotIn("query_responses", state)
        self.assertEqual(set(state["reviews"][0]), set(reviews.RUFUS_CSV_HEADERS))

    def test_public_docs_do_not_publish_internal_endpoint(self):
        root = Path(__file__).parents[1]
        public_files = [root / "README.md", root / "SKILL.md", *sorted((root / "references").glob("*.md"))]
        for path in public_files:
            text = path.read_text("utf-8")
            self.assertNotIn("/rufus/web/askSearch", text, path.name)

    def test_asins_are_validated_and_deduplicated(self):
        self.assertEqual(
            reviews.validate_asins(["b0gyn4835h", "B0H4YP63HV", "B0GYN4835H"]),
            ["B0GYN4835H", "B0H4YP63HV"],
        )
        with self.assertRaises(ValueError):
            reviews.validate_asins(["BAD-ASIN"])

    def test_parses_review_fields_and_date(self):
        html = """
        <ul><li>
          <span class="askReviewTitle">Safe seat</span>
          <span class="askReviewAuthor">Review by <a href="/gp/profile/P1">Amazon Customer</a> in the United States on August 16, 2026</span>
          <div class="askReviewText"><noscript></noscript><span>Works well.</span></div>
          <span role="img" aria-label="4.0 out of 5 stars"></span>
          <a href="/review/R-DATE-1">See full review</a>
        </li></ul>
        """
        rows = reviews.parse_rufus_html(
            html, "B000000001", source_url="https://www.amazon.com/rufus/web/askSearch",
            query="seat", capture_time="2026-08-24T02:00:00-07:00", http_status=200, batch=1,
        )
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["评论ID"], "R-DATE-1")
        self.assertEqual(rows[0]["评论人"], "Review by Amazon Customer in the United States on August 16, 2026")
        self.assertEqual(rows[0]["评论人链接"], "https://www.amazon.com/gp/profile/P1")
        self.assertEqual(rows[0]["评分"], 4.0)
        self.assertEqual(rows[0]["评论时间"], "2026-08-16")
        self.assertEqual(reviews.RUFUS_CSV_HEADERS, [
            "下载ASIN", "评论ID", "评论标题", "评论内容", "评论人", "评论人链接", "评分", "评论时间",
        ])

    def test_top_terms_remove_stopwords(self):
        rows = [
            {"评论标题": "The durable brush", "评论内容": "Durable durable and easy to use"},
            {"评论标题": "Quiet", "评论内容": "Quiet durable brush"},
        ]
        self.assertEqual(reviews.rufus_top_terms(rows), ["durable", "brush", "quiet", "easy"])

    def test_checkpoint_resume_and_restart(self):
        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp)
            reviews.atomic_write_json(output / "checkpoint.json", {
                "schema_version": 1,
                "asins": {"B000000001": {"status": "running", "completed_queries": ["quality"]}},
            })
            self.assertEqual(
                reviews.load_rufus_checkpoint(output)["asins"]["B000000001"]["completed_queries"],
                ["quality"],
            )
            self.assertEqual(reviews.load_rufus_checkpoint(output, restart=True), {"schema_version": 1, "asins": {}})

    def test_output_validation(self):
        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp)
            rows = [{header: "" for header in reviews.RUFUS_CSV_HEADERS}]
            rows[0].update({"下载ASIN": "B000000001", "评论ID": "R1", "评论内容": "Fixture"})
            workbook = reviews._make_workbook("Rufus评论", reviews.RUFUS_CSV_HEADERS, rows)
            reviews._save_workbook_atomic(workbook, output / "评论合并.xlsx")
            reviews.atomic_write_json(output / "B000000001_raw.json", {"reviews": rows})
            reviews.atomic_write_json(output / "汇总.json", {
                "asins": ["B000000001"], "counts": {"B000000001": 1}, "total": 1,
            })
            result = reviews.validate_output(output)
            self.assertTrue(result["ok"], result["errors"])
            workbook = load_workbook(output / "评论合并.xlsx", read_only=True)
            self.assertEqual(workbook.active.max_column, 8)
            workbook.close()

    def test_public_output_writer_preserves_columns_and_redacts_files(self):
        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp)
            row = {
                "下载ASIN": "B000000001", "评论ID": "R1", "评论标题": "Title", "评论内容": "Body",
                "评论人": "User", "评论人链接": "https://www.amazon.com/profile/P1", "评分": 5.0,
                "评论时间": "2026-08-24", "source_url": "https://internal.example/path", "query": "seat",
                "capture_time": "now", "http_status": 200, "batch": 1,
            }
            summary = reviews._write_public_outputs(
                output,
                ["B000000001"],
                {"B000000001": {"status": "complete", "stop_reason": "", "reviews": [row]}},
                "2026-08-24T00:00:00Z",
            )
            self.assertEqual(summary["total"], 1)
            workbook = load_workbook(output / "评论合并.xlsx", read_only=True, data_only=True)
            sheet = workbook.active
            self.assertEqual([cell.value for cell in next(sheet.iter_rows(max_row=1))], reviews.RUFUS_CSV_HEADERS)
            self.assertEqual(list(next(sheet.iter_rows(min_row=2, max_row=2, values_only=True))), [row[h] for h in reviews.RUFUS_CSV_HEADERS])
            workbook.close()
            raw = json.loads((output / "B000000001_raw.json").read_text("utf-8"))
            self.assertEqual(set(raw["reviews"][0]), set(reviews.RUFUS_CSV_HEADERS))
            public_text = "\n".join(
                path.read_text("utf-8", errors="ignore")
                for path in output.iterdir()
                if path.suffix in {".json", ".md", ".log", ".csv"}
            ).lower()
            for marker in ("source_url", "capture_time", "http_status", "query_responses", "profile_dir", "endpoint", "asksearch", "<html"):
                self.assertNotIn(marker, public_text)

    def test_profile_modes_and_limits(self):
        parser = reviews.build_parser()
        args = parser.parse_args(["rufus", "--asins", "B000000001", "--output", "out"])
        self.assertEqual(args.session_mode, "persistent")
        self.assertTrue(reviews._rufus_launch_plan(args)["persistent"])
        fresh = parser.parse_args([
            "rufus", "--asins", "B000000001", "--session-mode", "fresh", "--output", "out",
        ])
        self.assertFalse(reviews._rufus_launch_plan(fresh)["persistent"])
        reviews.validate_runtime_limits(4, 12)
        for workers, concurrency in ((0, 6), (5, 6), (2, 0), (2, 13)):
            with self.assertRaises(ValueError):
                reviews.validate_runtime_limits(workers, concurrency)

    def test_doctor_init_profile_and_validate_commands_exist(self):
        parser = reviews.build_parser()
        self.assertEqual(parser.parse_args(["doctor"]).command, "doctor")
        self.assertEqual(parser.parse_args(["init-profile"]).command, "init-profile")
        self.assertEqual(parser.parse_args(["validate", "--output", "out"]).command, "validate")
        self.assertNotIn("woot", parser.format_help())


class RufusBatchTests(unittest.IsolatedAsyncioTestCase):
    async def test_resume_batches_and_deduplicates(self):
        calls, checkpoints = [], []
        active = max_active = 0

        async def fetch_html(asin, query):
            nonlocal active, max_active
            calls.append(query)
            active += 1
            max_active = max(max_active, active)
            await asyncio.sleep(0.01)
            active -= 1
            review_id = "DUPLICATE" if query in {"q1", "q2"} else query.upper()
            return {
                "http_status": 200,
                "source_url": f"https://www.amazon.com/rufus/web/askSearch?asin={asin}&query={query}",
                "capture_time": "2026-08-24T02:00:00-07:00",
                "html": f'<li><a href="/review/{review_id}">review</a><span class="askReviewTitle">{query}</span><noscript></noscript>body</li>',
            }

        result = await reviews.run_rufus_queries(
            "B000000001", ["q0", "q1", "q2", "q3", "q4", "q5", "q6", "q7"],
            fetch_html=fetch_html, completed_queries={"q0"}, existing_reviews=[{"评论ID": "Q0"}],
            on_batch=checkpoints.append,
        )
        self.assertEqual(calls, ["q1", "q2", "q3", "q4", "q5", "q6", "q7"])
        self.assertEqual(max_active, 6)
        self.assertEqual(len(checkpoints), 2)
        self.assertEqual(len(result["reviews"]), 7)

    async def test_access_limits_stop_later_batches(self):
        calls = []

        async def forbidden(_asin, query):
            calls.append(query)
            return {"http_status": 403, "source_url": "https://fixture.test", "capture_time": "now", "html": "Forbidden"}

        result = await reviews.run_rufus_queries(
            "B000000001", ["first", "second"], fetch_html=forbidden, batch_size=1,
        )
        self.assertEqual(calls, ["first"])
        self.assertEqual(result["stop_reason"], "HTTP 403")


if __name__ == "__main__":
    unittest.main()
