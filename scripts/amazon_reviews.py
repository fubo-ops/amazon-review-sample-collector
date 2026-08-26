#!/usr/bin/env python3
"""Collect and validate Amazon Alexa for Shopping review samples."""

from __future__ import annotations

import argparse
import asyncio
import csv
import importlib.util
import json
import os
import platform
import re
import shutil
import sys
import tempfile
import time
from collections import Counter
from datetime import datetime
from html.parser import HTMLParser
from pathlib import Path
from typing import Any, Iterable
from urllib.parse import urlencode, urljoin

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


ASIN_RE = re.compile(r"^[A-Z0-9]{10}$")
MONTH_DATE_RE = re.compile(
    r"(January|February|March|April|May|June|July|August|September|October|November|December)\s+(\d{1,2}),\s+(\d{4})",
    re.IGNORECASE,
)
RUFUS_BASE_URL = "https://www.amazon.com/rufus/web/askSearch"
RUFUS_START_URL = "https://www.amazon.com/rd/uedata?"
SELLERSPRITE_REVERSE_URL = "https://www.sellersprite.com/v3/keyword-reverse/"
RUFUS_STOPWORDS = {
    "a", "about", "after", "all", "also", "am", "an", "and", "any", "are", "as", "at", "be", "because",
    "been", "but", "by", "can", "could", "did", "do", "does", "for", "from", "get", "got", "had", "has",
    "have", "he", "her", "hers", "him", "his", "how", "i", "if", "in", "into", "is", "it", "its", "just",
    "me", "more", "most", "my", "no", "not", "of", "on", "one", "or", "our", "out", "really", "she",
    "so", "some", "than", "that", "the", "their", "them", "then", "there", "these", "they", "this", "those",
    "to", "too", "up", "use", "used", "very", "was", "we", "were", "what", "when", "which", "who", "will",
    "with", "would", "you", "your",
}
EXPANSION_NOISE = {
    "amazon", "item", "items", "other", "others", "product", "products", "review", "reviews",
    "search", "term", "terms", "thing", "things", "para", "perro", "perros", "coche", "el",
}
SELLER_TERM_RE = re.compile(r"(?mi)^[ \t]*([a-z0-9][a-z0-9/-]*)\s+\(([\d,]+)[^)]*\)")


def default_profile_dir() -> Path:
    return Path.home() / ".amazon-review-collector" / "amazon-profile"


def validate_runtime_limits(asin_workers: int, concurrency: int) -> None:
    if not 1 <= asin_workers <= 4:
        raise ValueError("--asin-workers must be between 1 and 4")
    if not 1 <= concurrency <= 12:
        raise ValueError("--concurrency must be between 1 and 12")


def doctor_environment() -> dict[str, Any]:
    return {
        "python": platform.python_version(),
        "python_supported": (3, 10) <= sys.version_info[:2] <= (3, 13),
        "platform": platform.platform(),
        "playwright": importlib.util.find_spec("playwright") is not None,
        "openpyxl": importlib.util.find_spec("openpyxl") is not None,
    }


def configure_console() -> None:
    """Keep JSON and Chinese status text readable on Windows, macOS, and Linux."""
    for stream in (sys.stdout, sys.stderr):
        reconfigure = getattr(stream, "reconfigure", None)
        if reconfigure:
            reconfigure(encoding="utf-8", errors="replace")


def validate_asins(values: Iterable[str]) -> list[str]:
    result: list[str] = []
    seen: set[str] = set()
    for value in values:
        asin = value.strip().upper()
        if not ASIN_RE.fullmatch(asin):
            raise ValueError(f"Invalid ASIN: {value}")
        if asin not in seen:
            result.append(asin)
            seen.add(asin)
    if not result:
        raise ValueError("At least one ASIN is required")
    return result


def normalize_date(text: str) -> str:
    text = (text or "").strip()
    match = MONTH_DATE_RE.search(text)
    if match:
        return datetime.strptime(" ".join(match.groups()), "%B %d %Y").date().isoformat()
    iso_match = re.search(r"\b(\d{4}-\d{2}-\d{2})\b", text)
    return iso_match.group(1) if iso_match else text


def normalize_rating(text: Any) -> float | None:
    if isinstance(text, (int, float)):
        return float(text)
    match = re.search(r"(\d+(?:\.\d+)?)", str(text or ""))
    return float(match.group(1)) if match else None


def atomic_write_json(path: Path, value: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile("w", encoding="utf-8", suffix=".tmp", dir=path.parent, delete=False) as handle:
        json.dump(value, handle, ensure_ascii=False, indent=2, sort_keys=True)
        handle.write("\n")
        temp_path = Path(handle.name)
    os.replace(temp_path, path)


def load_rufus_checkpoint(output: Path, restart: bool = False) -> dict[str, Any]:
    path = Path(output) / "checkpoint.json"
    if restart or not path.exists():
        return {"schema_version": 1, "asins": {}}
    try:
        value = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        raise ValueError(f"Invalid checkpoint: {exc}") from exc
    if value.get("schema_version") != 1 or not isinstance(value.get("asins"), dict):
        raise ValueError("Invalid checkpoint schema")
    return value


def _save_workbook_atomic(workbook: Workbook, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temp_path = path.with_name(path.stem + ".tmp.xlsx")
    workbook.save(temp_path)
    load_workbook(temp_path, read_only=True).close()
    os.replace(temp_path, path)


def _make_workbook(sheet_name: str, headers: list[str], rows: Iterable[dict[str, Any]]) -> Workbook:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_name
    sheet.append(headers)
    for row in rows:
        sheet.append([row.get(header, "") for header in headers])
    fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    widths = {
        "评论ID": 30, "下载ASIN": 14, "评论ASIN": 14, "评论时间": 14,
        "评论人": 22, "评分": 10, "评论": 70, "评论人链接": 42,
        "评论标题": 40, "评论内容": 70, "好评差评": 12, "评论国家": 20,
    }
    for index, header in enumerate(headers, 1):
        sheet.column_dimensions[get_column_letter(index)].width = widths.get(header, 16)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    return workbook


def validate_output(output: Path) -> dict[str, Any]:
    output = Path(output)
    summary_path = output / "汇总.json"
    if not summary_path.exists():
        return {"ok": False, "errors": ["汇总.json is missing"]}

    errors: list[str] = []
    summary = json.loads(summary_path.read_text("utf-8"))
    merged_path = output / "评论合并.xlsx"
    workbook_count = -1
    if not merged_path.exists():
        errors.append("Rufus merged workbook missing")
    else:
        workbook = load_workbook(merged_path, read_only=True, data_only=True)
        sheet = workbook.active
        headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
        if headers != RUFUS_CSV_HEADERS:
            errors.append("Rufus workbook headers mismatch")
        workbook_count = max(sheet.max_row - 1, 0)
        workbook.close()

    raw_total = 0
    for asin in summary.get("asins", []):
        raw_path = output / f"{asin}_raw.json"
        if not raw_path.exists():
            errors.append(f"Rufus raw file missing: {asin}")
            continue
        rows = json.loads(raw_path.read_text("utf-8")).get("reviews", [])
        ids = [str(row.get("评论ID") or "") for row in rows]
        if any(not review_id for review_id in ids):
            errors.append(f"empty review ID: {asin}")
        if len(ids) != len(set(ids)):
            errors.append(f"duplicate review ID: {asin}")
        if len(rows) != int(summary.get("counts", {}).get(asin, -1)):
            errors.append(f"Rufus raw count mismatch: {asin}")
        raw_total += len(rows)

    if raw_total != int(summary.get("total", -1)) or workbook_count != raw_total:
        errors.append("Rufus merged row count mismatch")
    return {"ok": not errors, "errors": errors, "summary": summary}

class _HtmlNode:
    def __init__(self, tag: str, attrs: list[tuple[str, str | None]], parent: "_HtmlNode | None" = None):
        self.tag = tag
        self.attrs = {key: value or "" for key, value in attrs}
        self.parent = parent
        self.children: list[_HtmlNode | str] = []

    def text(self) -> str:
        return " ".join(
            part for child in self.children if (part := child.text() if isinstance(child, _HtmlNode) else child.strip())
        ).strip()


class _RufusHtmlParser(HTMLParser):
    VOID_TAGS = {"area", "base", "br", "col", "embed", "hr", "img", "input", "link", "meta", "param", "source", "track", "wbr"}

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.root = _HtmlNode("document", [])
        self.stack = [self.root]

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        node = _HtmlNode(tag.lower(), attrs, self.stack[-1])
        self.stack[-1].children.append(node)
        if tag.lower() not in self.VOID_TAGS:
            self.stack.append(node)

    def handle_startendtag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        self.stack[-1].children.append(_HtmlNode(tag.lower(), attrs, self.stack[-1]))

    def handle_endtag(self, tag: str) -> None:
        tag = tag.lower()
        for index in range(len(self.stack) - 1, 0, -1):
            if self.stack[index].tag == tag:
                del self.stack[index:]
                break

    def handle_data(self, data: str) -> None:
        self.stack[-1].children.append(data)


def _walk_html(node: _HtmlNode) -> Iterable[_HtmlNode]:
    yield node
    for child in node.children:
        if isinstance(child, _HtmlNode):
            yield from _walk_html(child)


def _node_with_class(node: _HtmlNode, class_name: str) -> _HtmlNode | None:
    return next((item for item in _walk_html(node) if class_name in item.attrs.get("class", "").split()), None)


def _body_after_noscript(li: _HtmlNode) -> str:
    noscript = next((item for item in _walk_html(li) if item.tag == "noscript"), None)
    if noscript is None or noscript.parent is None:
        return ""
    siblings = noscript.parent.children
    for child in siblings[siblings.index(noscript) + 1:]:
        text = child.text() if isinstance(child, _HtmlNode) else child.strip()
        if text:
            return re.sub(r"\s+", " ", text).strip()
    return ""


def parse_rufus_html(
    html: str,
    asin: str,
    *,
    source_url: str,
    query: str,
    capture_time: str,
    http_status: int,
    batch: int,
) -> list[dict[str, Any]]:
    parser = _RufusHtmlParser()
    parser.feed(html or "")
    rows: list[dict[str, Any]] = []
    for li in (node for node in _walk_html(parser.root) if node.tag == "li"):
        review_link = next(
            (node for node in _walk_html(li) if node.tag == "a" and "/review/" in node.attrs.get("href", "")),
            None,
        )
        if review_link is None:
            continue
        match = re.search(r"/review/([^/?#]+)", review_link.attrs.get("href", ""))
        if not match:
            continue
        title_node = _node_with_class(li, "askReviewTitle")
        author_node = _node_with_class(li, "askReviewAuthor")
        profile = next(
            (node.attrs.get("href", "") for node in _walk_html(author_node or li) if node.tag == "a" and "/profile/" in node.attrs.get("href", "")),
            "",
        )
        rating_node = next((node for node in _walk_html(li) if node.attrs.get("role") == "img" and node.attrs.get("aria-label")), None)
        full_text = re.sub(r"\s+", " ", li.text())
        date_match = MONTH_DATE_RE.search(full_text)
        date_text = date_match.group(0) if date_match else ""
        rows.append({
            "下载ASIN": asin,
            "评论ID": match.group(1),
            "评论标题": title_node.text() if title_node else "",
            "评论内容": _body_after_noscript(li),
            "评论人": author_node.text() if author_node else "",
            "评论人链接": urljoin("https://www.amazon.com", profile) if profile else "",
            "评分": normalize_rating(rating_node.attrs.get("aria-label", "") if rating_node else ""),
            "评论时间": normalize_date(date_text),
            "source_url": source_url,
            "query": query,
            "capture_time": capture_time,
            "http_status": int(http_status),
            "batch": int(batch),
        })
    return rows


def rufus_top_terms(
    rows: Iterable[dict[str, Any]],
    limit: int = 30,
    excluded: Iterable[str] = (),
) -> list[str]:
    blocked = RUFUS_STOPWORDS | EXPANSION_NOISE | {word.lower() for word in excluded}
    counts: Counter[str] = Counter()
    for row in rows:
        text = f"{row.get('评论标题', '')} {row.get('评论内容', '')}".lower()
        counts.update(word for word in re.findall(r"[a-z][a-z'-]*", text) if word not in blocked and len(word) > 1)
    return [word for word, _count in sorted(counts.items(), key=lambda item: (-item[1], item[0]))[:limit]]


def select_seller_queries(text: str, excluded: Iterable[str] = (), limit: int = 30) -> list[str]:
    starts = [position for marker in ("高频词", "Word Frequency") if (position := text.find(marker)) >= 0]
    start = min(starts) if starts else -1
    if start < 0:
        return []
    ends = [
        position for marker in ("\n对比", "\nAdd All to My List", "\nExport Keywords")
        if (position := text.find(marker, start)) > start
    ]
    section = text[start:min(ends) if ends else None]
    blocked = EXPANSION_NOISE | {word.lower() for word in excluded}
    result: list[str] = []
    seen: set[str] = set()
    for match in SELLER_TERM_RE.finditer(section):
        word = match.group(1).lower()
        if word not in blocked and word not in seen:
            result.append(word)
            seen.add(word)
        if len(result) >= limit:
            break
    return result


def _rufus_stop_reason(status: int, html: str, url: str = "") -> str:
    if status in {403, 429, 503}:
        return f"HTTP {status}"
    combined = f"{html} {url}".lower()
    if any(marker in combined for marker in ("captcha", "robot check", "not a robot", "awswaf", "validatecaptcha", "access denied", "too many requests", "rate limit")):
        return "CAPTCHA or access-limit response"
    if "/ap/signin" in url.lower() or "sign in to see search results and to post your question" in combined:
        return "login response"
    if status != 200:
        return f"HTTP {status}" if status else "network error"
    return ""


async def run_rufus_queries(
    asin: str,
    queries: list[str],
    *,
    fetch_html: Any,
    batch_size: int = 6,
    completed_queries: set[str] | None = None,
    existing_reviews: Iterable[dict[str, Any]] = (),
    on_batch: Any = None,
) -> dict[str, Any]:
    completed = set(completed_queries or ())
    pending = [query for query in queries if query not in completed]
    result: dict[str, Any] = {
        "reviews": [], "responses": [], "completed_queries": [], "stopped": False, "stop_reason": "",
    }
    unique = {str(row["评论ID"]): row for row in existing_reviews}
    for offset in range(0, len(pending), batch_size):
        batch_number = offset // batch_size + 1
        batch_queries = pending[offset:offset + batch_size]
        fetched = await asyncio.gather(*(fetch_html(asin, query) for query in batch_queries))
        stop_reason = ""
        for query, response in zip(batch_queries, fetched):
            evidence = {**response, "query": query, "batch": batch_number}
            result["responses"].append(evidence)
            reason = _rufus_stop_reason(int(response.get("http_status") or 0), str(response.get("html") or ""), str(response.get("source_url") or ""))
            if reason:
                stop_reason = stop_reason or reason
                continue
            completed.add(query)
            rows = parse_rufus_html(
                str(response.get("html") or ""),
                asin,
                source_url=str(response.get("source_url") or ""),
                query=query,
                capture_time=str(response.get("capture_time") or ""),
                http_status=int(response.get("http_status") or 0),
                batch=batch_number,
            )
            for row in rows:
                unique.setdefault(row["评论ID"], row)
        result.update({
            "reviews": list(unique.values()),
            "completed_queries": [query for query in queries if query in completed],
        })
        if on_batch:
            on_batch({**result, "responses": list(result["responses"])})
        if stop_reason:
            result.update({"stopped": True, "stop_reason": stop_reason})
            return result
    result.update({
        "reviews": list(unique.values()),
        "completed_queries": [query for query in queries if query in completed],
    })
    return result


def _ordered_unique(values: Iterable[str], included: set[str] | None = None) -> list[str]:
    allowed = included if included is not None else set(values)
    result: list[str] = []
    seen: set[str] = set()
    for value in values:
        if value in allowed and value not in seen:
            result.append(value)
            seen.add(value)
    return result


async def run_hybrid_expansion(
    asin: str,
    *,
    first_queries: list[str],
    first_reviews: Iterable[dict[str, Any]],
    seller_queries: list[str],
    fetch_html: Any,
    second_limit: int = 30,
    batch_size: int = 6,
    completed_queries: set[str] | None = None,
    second_queries: list[str] | None = None,
    on_batch: Any = None,
) -> dict[str, Any]:
    """Expand an existing sample with seller terms, then terms found in the enlarged sample."""
    completed = set(completed_queries or ()) | set(first_queries)
    reviews = list(first_reviews)
    seller_terms = [term for term in seller_queries if term not in completed]

    def notify(phase: str, snapshot: dict[str, Any], derived: list[str]) -> None:
        if on_batch:
            snapshot_completed = completed | set(snapshot.get("completed_queries") or [])
            on_batch({
                **snapshot,
                "phase": phase,
                "completed_queries": _ordered_unique(
                    [*first_queries, *seller_queries, *derived], snapshot_completed
                ),
                "seller_queries": seller_queries,
                "second_queries": derived,
            })

    if seller_terms:
        seller_result = await run_rufus_queries(
            asin,
            seller_terms,
            fetch_html=fetch_html,
            batch_size=batch_size,
            completed_queries=completed,
            existing_reviews=reviews,
            on_batch=lambda snapshot: notify("seller", snapshot, []),
        )
        reviews = seller_result["reviews"]
        completed.update(seller_result["completed_queries"])
        if seller_result["stopped"]:
            return {
                "reviews": reviews,
                "completed_queries": _ordered_unique([*first_queries, *seller_queries], completed),
                "seller_queries": seller_queries,
                "second_queries": [],
                "status": "blocked",
                "stopped": True,
                "stop_reason": seller_result["stop_reason"],
            }

    if second_queries is None:
        second_queries = rufus_top_terms(
            reviews,
            limit=second_limit,
            excluded=completed | set(seller_queries),
        )
    else:
        second_queries = list(second_queries)
    second_terms = [term for term in second_queries if term not in completed]
    second_result = await run_rufus_queries(
        asin,
        second_terms,
        fetch_html=fetch_html,
        batch_size=batch_size,
        completed_queries=completed,
        existing_reviews=reviews,
        on_batch=lambda snapshot: notify("second", snapshot, second_queries or []),
    )
    reviews = second_result["reviews"]
    completed.update(second_result["completed_queries"])
    stopped = bool(second_result["stopped"])
    return {
        "reviews": reviews,
        "completed_queries": _ordered_unique(
            [*first_queries, *seller_queries, *(second_queries or [])], completed
        ),
        "seller_queries": seller_queries,
        "second_queries": second_queries or [],
        "status": "blocked" if stopped else "complete",
        "stopped": stopped,
        "stop_reason": second_result["stop_reason"],
    }


async def collect_seller_queries(
    page: Any,
    asin: str,
    *,
    excluded: Iterable[str] = (),
    limit: int = 30,
    timeout_ms: int = 20000,
) -> list[str]:
    """Read reverse-search terms from the signed-in seller research page."""
    url = f"{SELLERSPRITE_REVERSE_URL}?{urlencode({'q': asin, 'marketId': 1})}"
    try:
        await page.goto(url, wait_until="domcontentloaded", timeout=timeout_ms)
        button = page.get_by_role("button", name="立即查询", exact=True)
        if await button.count():
            await button.first.click(timeout=min(timeout_ms, 5000))
        deadline = time.monotonic() + timeout_ms / 1000
        last_terms: list[str] = []
        while time.monotonic() < deadline:
            text = await page.locator("body").inner_text(timeout=min(timeout_ms, 5000))
            last_terms = select_seller_queries(text, excluded=excluded, limit=limit)
            if last_terms:
                return last_terms
            await asyncio.sleep(1)
        return last_terms
    except Exception:
        return []


async def _rufus_page_fetch(page: Any, asin: str, query: str, timeout_ms: int) -> dict[str, Any]:
    params = {
        "isBlue": "true",
        "asin": asin,
        "query": query,
        "abar_qid": int(time.time() * 1000),
    }
    url = f"{RUFUS_BASE_URL}?{urlencode(params)}"
    capture_time = datetime.now().astimezone().isoformat(timespec="seconds")
    try:
        response = await page.evaluate(
            """
            async ({url, timeoutMs}) => {
              const controller = new AbortController();
              const timer = setTimeout(() => controller.abort(), timeoutMs);
              try {
                const response = await fetch(url, {method: 'GET', credentials: 'same-origin', signal: controller.signal});
                return {http_status: response.status, source_url: response.url || url, html: await response.text()};
              } finally {
                clearTimeout(timer);
              }
            }
            """,
            {"url": url, "timeoutMs": timeout_ms},
        )
        return {**response, "capture_time": capture_time}
    except Exception as exc:
        return {"http_status": 0, "source_url": url, "capture_time": capture_time, "html": "", "error": f"{type(exc).__name__}: {exc}"}


RUFUS_CSV_HEADERS = [
    "下载ASIN", "评论ID", "评论标题", "评论内容", "评论人", "评论人链接", "评分", "评论时间",
]


def _public_review(row: dict[str, Any]) -> dict[str, Any]:
    return {field: row.get(field, "") for field in RUFUS_CSV_HEADERS}


def _public_reason(reason: str) -> str:
    value = str(reason or "").lower()
    if not value:
        return ""
    if "login" in value or "sign" in value:
        return "需要重新登录"
    if "captcha" in value or "robot" in value:
        return "需要人工验证"
    if any(code in value for code in ("403", "429", "503", "access-limit")):
        return "访问暂时受限"
    if "timeout" in value:
        return "请求超时"
    return "采集未完成"


def _checkpoint_state(state: dict[str, Any]) -> dict[str, Any]:
    """Persist only the minimum restart state; never persist URLs, HTML, paths, or HTTP details."""
    return {
        "asin": str(state.get("asin") or ""),
        "status": str(state.get("status") or "running"),
        "stop_reason": _public_reason(str(state.get("stop_reason") or "")),
        "seed_complete": bool(state.get("seed_complete")),
        "expansion": str(state.get("expansion") or "hybrid"),
        "top30": list(state.get("top30") or []),
        "seller_queries": list(state.get("seller_queries") or []),
        "second_queries": list(state.get("second_queries") or []),
        "completed_queries": list(state.get("completed_queries") or []),
        "reviews": [_public_review(row) for row in state.get("reviews") or []],
    }


def _non_empty_rates(rows: list[dict[str, Any]]) -> dict[str, float]:
    if not rows:
        return {field: 0.0 for field in RUFUS_CSV_HEADERS}
    return {
        field: round(sum(row.get(field) not in (None, "") for row in rows) / len(rows), 4)
        for field in RUFUS_CSV_HEADERS
    }


def _write_rufus_report(output: Path, summary: dict[str, Any]) -> None:
    lines = [
        "# Amazon 评论样本采集说明",
        "",
        "数据来自 Amazon 商品评论相关结果，输出前已去重并移除运行参数、会话路径、请求地址和原始响应等技术信息。",
        "该结果是关键词召回的评论样本，不代表该 ASIN 的全部评论，也不适合推算全量评论的精确占比。",
        "",
        "| ASIN | 唯一评论数 | 状态 | 停止原因 |",
        "|---|---:|---|---|",
    ]
    for asin in summary["asins"]:
        lines.append(
            f"| {asin} | {summary['counts'][asin]} | {summary['statuses'][asin]} | "
            f"{summary.get('stop_reasons', {}).get(asin, '')} |"
        )
    lines.extend([
        "",
        "## 覆盖限制",
        "",
        "Amazon 相关性结果窗口有限。未被召回、超出结果窗口或 Amazon 未返回的内容都会遗漏。",
    ])
    (output / "采集说明.md").write_text("\n".join(lines) + "\n", encoding="utf-8")


def _write_public_outputs(
    output: Path,
    asins: list[str],
    results: dict[str, dict[str, Any]],
    started_at: str,
) -> dict[str, Any]:
    for asin, result in results.items():
        atomic_write_json(output / f"{asin}_raw.json", {
            "asin": asin,
            "status": result["status"],
            "stop_reason": _public_reason(result["stop_reason"]),
            "reviews": [_public_review(row) for row in result["reviews"]],
        })

    all_rows = [row for asin in asins for row in results[asin]["reviews"]]
    with (output / "评论合并.csv").open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=RUFUS_CSV_HEADERS)
        writer.writeheader()
        writer.writerows(_public_review(row) for row in all_rows)
    _save_workbook_atomic(
        _make_workbook("Rufus评论", RUFUS_CSV_HEADERS, all_rows),
        output / "评论合并.xlsx",
    )

    summary = {
        "coverage": "keyword-retrieved review sample; not all Amazon reviews",
        "asins": asins,
        "statuses": {asin: results[asin]["status"] for asin in asins},
        "stop_reasons": {asin: _public_reason(results[asin]["stop_reason"]) for asin in asins},
        "counts": {asin: len(results[asin]["reviews"]) for asin in asins},
        "unique_ids": {asin: len({row["评论ID"] for row in results[asin]["reviews"]}) for asin in asins},
        "field_non_empty_rate": {asin: _non_empty_rates(results[asin]["reviews"]) for asin in asins},
        "total": len(all_rows),
        "generated_at": datetime.now().astimezone().isoformat(timespec="seconds"),
    }
    atomic_write_json(output / "汇总.json", summary)
    log_lines = [f"{started_at} start"]
    for asin in asins:
        result = results[asin]
        log_lines.append(
            f"{datetime.now().astimezone().isoformat(timespec='seconds')} asin={asin} status={result['status']} "
            f"reviews={len(result['reviews'])} reason={_public_reason(result['stop_reason'])!r}"
        )
    (output / "采集日志.log").write_text("\n".join(log_lines) + "\n", encoding="utf-8")
    _write_rufus_report(output, summary)
    return summary


def _rufus_launch_plan(args: argparse.Namespace) -> dict[str, Any]:
    mode = getattr(args, "session_mode", "persistent")
    launch: dict[str, Any] = {"headless": not getattr(args, "headed", False)}
    if args.browser == "chrome":
        launch["channel"] = "chrome"
    if mode == "persistent":
        profile_dir = Path(getattr(args, "profile_dir", None) or default_profile_dir())
        return {"persistent": True, "profile_dir": str(profile_dir.expanduser().resolve()), "launch": launch}
    return {"persistent": False, "profile_dir": None, "launch": launch}


async def init_profile(args: argparse.Namespace) -> dict[str, Any]:
    from playwright.async_api import async_playwright

    profile_dir = Path(args.profile_dir).expanduser().resolve()
    profile_dir.mkdir(parents=True, exist_ok=True)
    launch: dict[str, Any] = {"headless": False}
    if args.browser == "chrome":
        launch["channel"] = "chrome"
    async with async_playwright() as playwright:
        context = await playwright.chromium.launch_persistent_context(
            str(profile_dir), locale="en-US", **launch
        )
        amazon_page = context.pages[0] if context.pages else await context.new_page()
        seller_page = await context.new_page()
        await amazon_page.goto("https://www.amazon.com/", wait_until="domcontentloaded", timeout=args.timeout_ms)
        await seller_page.goto(SELLERSPRITE_REVERSE_URL, wait_until="domcontentloaded", timeout=args.timeout_ms)
        print("Amazon 和卖家精灵已分别打开。请在两个标签页完成登录，然后回到终端按 Enter 保存此专用 Profile。")
        await asyncio.to_thread(input)
        await context.close()
    return {"status": "initialized"}


async def collect_rufus_reviews(args: argparse.Namespace) -> dict[str, Any]:
    from playwright.async_api import async_playwright

    asins = validate_asins(args.asins)
    validate_runtime_limits(args.asin_workers, args.concurrency)
    output = Path(args.output).expanduser().resolve()
    output.mkdir(parents=True, exist_ok=True)
    legacy_evidence = output / "响应证据"
    if legacy_evidence.exists():
        shutil.rmtree(legacy_evidence)
    checkpoint = load_rufus_checkpoint(output, restart=args.restart)
    checkpoint_path = output / "checkpoint.json"
    atomic_write_json(checkpoint_path, checkpoint)
    launch_plan = _rufus_launch_plan(args)
    started_at = datetime.now().astimezone().isoformat(timespec="seconds")

    async with async_playwright() as playwright:
        browser = None
        if launch_plan["persistent"]:
            Path(launch_plan["profile_dir"]).mkdir(parents=True, exist_ok=True)
            context = await playwright.chromium.launch_persistent_context(
                launch_plan["profile_dir"], locale="en-US", **launch_plan["launch"]
            )
        else:
            browser = await playwright.chromium.launch(**launch_plan["launch"])
            context = await browser.new_context(locale="en-US")
        async def route_handler(route: Any) -> None:
            if route.request.resource_type in {"image", "font", "media"}:
                await route.abort()
            else:
                await route.continue_()

        await context.route("**/*", route_handler)
        request_gate = asyncio.Semaphore(args.concurrency)
        asin_gate = asyncio.Semaphore(args.asin_workers)

        def save_state(asin: str, state: dict[str, Any]) -> None:
            checkpoint["asins"][asin] = _checkpoint_state(state)
            atomic_write_json(checkpoint_path, checkpoint)

        async def collect_asin(asin: str) -> tuple[str, dict[str, Any]]:
            saved = checkpoint["asins"].get(asin, {})
            if saved and saved.get("expansion", "baseline") != args.expansion:
                saved = {}
            if saved.get("status") == "complete":
                return asin, saved

            async with asin_gate:
                page = await context.new_page()
                try:
                    try:
                        response = await page.goto(RUFUS_START_URL, wait_until="domcontentloaded", timeout=args.timeout_ms)
                        start_status = response.status if response else 0
                        start_html = await page.content()
                        start_url = page.url
                    except Exception as exc:
                        start_status, start_html, start_url = 0, f"{type(exc).__name__}: {exc}", page.url
                    start_reason = _rufus_stop_reason(start_status, start_html, start_url)
                    if start_reason:
                        result = {
                            "asin": asin, "seed_complete": False, "top30": [],
                            "completed_queries": [], "reviews": [], "stopped": True,
                            "stop_reason": f"startup: {start_reason}", "status": "blocked",
                        }
                        save_state(asin, result)
                        return asin, result

                    async def fetch_html(_asin: str, query: str) -> dict[str, Any]:
                        async with request_gate:
                            return await _rufus_page_fetch(page, _asin, query, args.timeout_ms)

                    seed = None
                    seed_rows: list[dict[str, Any]] = []
                    seed_reason = ""
                    if saved.get("seed_complete"):
                        terms = list(saved.get("top30") or [])
                    else:
                        seed = await fetch_html(asin, "")
                        seed.update({"query": "", "batch": 0})
                        seed_reason = _rufus_stop_reason(
                            int(seed.get("http_status") or 0), str(seed.get("html") or ""), str(seed.get("source_url") or "")
                        )
                        seed_rows = [] if seed_reason else parse_rufus_html(
                            str(seed.get("html") or ""), asin, source_url=str(seed.get("source_url") or ""), query="",
                            capture_time=str(seed.get("capture_time") or ""), http_status=int(seed.get("http_status") or 0), batch=0,
                        )
                        terms = rufus_top_terms(seed_rows, 30)
                    resumed_reviews = {row["评论ID"]: row for row in seed_rows}
                    for row in saved.get("reviews") or []:
                        resumed_reviews.setdefault(row["评论ID"], row)
                    base_state = {
                        "asin": asin, "seed_complete": not bool(seed_reason), "top30": terms,
                        "expansion": args.expansion,
                        "seller_queries": list(saved.get("seller_queries") or []),
                        "second_queries": list(saved.get("second_queries") or []),
                        "completed_queries": list(saved.get("completed_queries") or []),
                        "reviews": list(resumed_reviews.values()), "stopped": bool(seed_reason),
                        "stop_reason": seed_reason, "status": "blocked" if seed_reason else "running",
                    }
                    save_state(asin, base_state)
                    if seed_reason:
                        return asin, base_state

                    def checkpoint_batch(snapshot: dict[str, Any]) -> None:
                        state = {
                            **base_state,
                            "completed_queries": snapshot["completed_queries"],
                            "reviews": snapshot["reviews"],
                            "stopped": snapshot["stopped"],
                            "stop_reason": snapshot["stop_reason"],
                            "status": "blocked" if snapshot["stopped"] else "running",
                        }
                        save_state(asin, state)

                    query_result = await run_rufus_queries(
                        asin,
                        terms,
                        fetch_html=fetch_html,
                        batch_size=min(6, args.concurrency),
                        completed_queries=set(base_state["completed_queries"]),
                        existing_reviews=base_state["reviews"],
                        on_batch=checkpoint_batch,
                    )
                    first_completed = set(base_state["completed_queries"]) | set(query_result["completed_queries"])
                    stop_reason = query_result["stop_reason"]
                    if stop_reason or args.expansion == "baseline":
                        final = {
                            **base_state,
                            "completed_queries": _ordered_unique(
                                [*base_state["completed_queries"], *terms], first_completed
                            ),
                            "reviews": query_result["reviews"],
                            "stopped": bool(stop_reason),
                            "stop_reason": stop_reason,
                            "status": "blocked" if stop_reason else "complete",
                        }
                        save_state(asin, final)
                        return asin, final

                    seller_terms = list(saved.get("seller_queries") or [])
                    if args.expansion == "hybrid" and not seller_terms:
                        seller_page = await context.new_page()
                        try:
                            seller_terms = await collect_seller_queries(
                                seller_page,
                                asin,
                                excluded=first_completed,
                                limit=args.seller_term_limit,
                                timeout_ms=args.seller_timeout_ms,
                            )
                        finally:
                            await seller_page.close()

                    second_saved = list(saved.get("second_queries") or []) or None

                    def checkpoint_expansion(snapshot: dict[str, Any]) -> None:
                        save_state(asin, {
                            **base_state,
                            "seller_queries": snapshot.get("seller_queries") or seller_terms,
                            "second_queries": snapshot.get("second_queries") or [],
                            "completed_queries": snapshot["completed_queries"],
                            "reviews": snapshot["reviews"],
                            "stopped": snapshot["stopped"],
                            "stop_reason": snapshot["stop_reason"],
                            "status": "blocked" if snapshot["stopped"] else "running",
                        })

                    expanded = await run_hybrid_expansion(
                        asin,
                        first_queries=terms,
                        first_reviews=query_result["reviews"],
                        seller_queries=seller_terms,
                        fetch_html=fetch_html,
                        second_limit=args.second_term_limit,
                        batch_size=min(6, args.concurrency),
                        completed_queries=first_completed,
                        second_queries=second_saved,
                        on_batch=checkpoint_expansion,
                    )
                    final = {
                        **base_state,
                        "seller_queries": expanded["seller_queries"],
                        "second_queries": expanded["second_queries"],
                        "completed_queries": expanded["completed_queries"],
                        "reviews": expanded["reviews"],
                        "stopped": expanded["stopped"],
                        "stop_reason": expanded["stop_reason"],
                        "status": expanded["status"],
                    }
                    save_state(asin, final)
                    return asin, final
                finally:
                    await page.close()

        try:
            pairs = await asyncio.gather(*(collect_asin(asin) for asin in asins))
            results = dict(pairs)
        finally:
            await context.close()
            if browser is not None:
                await browser.close()

    summary = _write_public_outputs(output, asins, results, started_at)
    if all(status == "complete" for status in summary["statuses"].values()):
        checkpoint_path.unlink(missing_ok=True)
    return summary


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description=__doc__)
    subparsers = parser.add_subparsers(dest="command", required=True)

    rufus = subparsers.add_parser("rufus", help="collect an Amazon Alexa for Shopping review sample")
    rufus.add_argument("--asins", nargs="+", required=True)
    rufus.add_argument("--browser", choices=("chromium", "chrome"), default="chromium")
    rufus.add_argument("--session-mode", choices=("fresh", "persistent"), default="persistent")
    rufus.add_argument("--profile-dir", default=str(default_profile_dir()))
    rufus.add_argument("--headed", action="store_true")
    rufus.add_argument("--asin-workers", type=int, default=2)
    rufus.add_argument("--concurrency", type=int, default=6)
    rufus.add_argument("--expansion", choices=("hybrid", "second-pass", "baseline"), default="hybrid")
    rufus.add_argument("--seller-term-limit", type=int, default=30)
    rufus.add_argument("--second-term-limit", type=int, default=30)
    rufus.add_argument("--seller-timeout-ms", type=int, default=20000)
    rufus.add_argument("--restart", action="store_true")
    rufus.add_argument("--timeout-ms", type=int, default=30000)
    rufus.add_argument("--output", required=True)

    init = subparsers.add_parser("init-profile", help="open the dedicated Amazon profile for sign-in")
    init.add_argument("--profile-dir", default=str(default_profile_dir()))
    init.add_argument("--browser", choices=("chromium", "chrome"), default="chromium")
    init.add_argument("--timeout-ms", type=int, default=30000)
    subparsers.add_parser("doctor", help="report runtime readiness")
    validate = subparsers.add_parser("validate", help="validate an output directory")
    validate.add_argument("--output", required=True)
    return parser

def main(argv: list[str] | None = None) -> int:
    configure_console()
    args = build_parser().parse_args(argv)
    try:
        if args.command == "rufus":
            validate_runtime_limits(args.asin_workers, args.concurrency)
            result = asyncio.run(collect_rufus_reviews(args))
            print(json.dumps(result, ensure_ascii=False, indent=2))
            return 2 if any(status == "blocked" for status in result["statuses"].values()) else 0
        if args.command == "init-profile":
            result = asyncio.run(init_profile(args))
            print(json.dumps(result, ensure_ascii=False, indent=2))
            return 0
        if args.command == "doctor":
            result = doctor_environment()
            print(json.dumps(result, ensure_ascii=False, indent=2))
            return 0 if result["python_supported"] and result["playwright"] and result["openpyxl"] else 1
        result = validate_output(Path(args.output))
        print(json.dumps(result, ensure_ascii=False, indent=2))
        return 0 if result.get("ok") else 1
    except (ValueError, RuntimeError) as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 1

if __name__ == "__main__":
    raise SystemExit(main())
